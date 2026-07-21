import openpyxl
import random
from PIL import Image
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import cm_to_EMU
import subprocess
import os
from pathlib import Path
# ================= НАСТРОЙКИ =================
# Путь к исходному файлу (должен существовать)
input_file = 'data_base_py/best_power_protocol.xlsx'
# Папка, в которую нужно сохранить результат
# output_folder = 'posle_gotovie_protocol'
# =============================================

def input_rename_new_protocol(output_folder, inventory_protocol, mux_protocol, transmitter_protocol,
                              izgotovitel_protocol,
                              serial_number_protocol, cell_id_protocol, rich_protocol, tvk_protocol, date_protocol,
                              name_naselennogo_puncta, chastota, model_protocol, replace_text_power, replace_text_MER,
                              replace_text_neravnomernost_achh, replace_text_frequency_offset,location_itog):
    # Формируем полный путь для сохранения
    # Имя нового файла
    # output_filename = f"Протокол_зоны_НЦТВ_{itog_rayon}_{itog_location_measure_metrics}_{date_protocol}.xlsx"
    output_filename = f"Протокол_{date_protocol}_{name_naselennogo_puncta}_{mux_protocol}_{transmitter_protocol}_{inventory_protocol}.xlsx"
    # output_filename = f"Протокол_{clean_date}_{location}_{mux}_{transmitter}_{clean_inventory}.pdf"
    output_path = os.path.join(output_folder, output_filename)

    # 1. Создаем папку для сохранения, если она еще не существует
    os.makedirs(output_folder, exist_ok=True)

    # 2. Открываем существующий Excel файл
    try:
        workbook = openpyxl.load_workbook(input_file)
    except FileNotFoundError:
        print(f"Ошибка: Файл '{input_file}' не найден!")
        exit()

        # Выбираем лист для редактирования
        # (можно указать имя листа: workbook['Лист1'] или оставить workbook.active для текущего)
    sheet = workbook.active

    # 3. Вставляем данные
    # --- Вариант А: Вставка в конкретные ячейки по координатам ---
    sheet['E5'] = date_protocol
    sheet['H31'] = date_protocol
    sheet['I11'] = location_itog
    sheet['I12'] = str(izgotovitel_protocol) + " " + str(model_protocol)
    sheet['I13'] = str(tvk_protocol) + "(ТВК), " + str(chastota) + " (МГц)"
    sheet['I14'] = inventory_protocol
    sheet['I15'] = serial_number_protocol
    sheet['O21'] = 3 - int(mux_protocol)
    sheet['O22'] = cell_id_protocol
    sheet['I48'] = replace_text_power
    sheet['I49'] = "0.0E-09"
    sheet['I50'] = replace_text_MER
    sheet['I51'] = replace_text_neravnomernost_achh
    sheet['I53'] = round(random.randint(-900, 900) / 10, 1)
    sheet['I54'] = replace_text_frequency_offset
    sheet['N61'] = "Авдяков Е.А." if int(date_protocol.split('-')[-1]) % 2 == 0 else "Халатаев Т.С."
    sheet['H66'] = location_itog
    # 4. Сохраняем файл по новому пути
    workbook.save(output_path)
    # print(f"✅ Файл успешно сохранен по пути:\n{output_path}")


def insert_png_folder_to_xlsx(
        xlsx_path: str,
        image_folder: str,
        sheet_name: str = "title",
        # Конфигурация каждой картинки по имени файла
        images_config: dict[str, dict] | None = None,
        # Параметры по умолчанию для картинок без явной конфигурации
        default_width_cm: float = 15.0,
        default_height_cm: float = 10.0,
        default_x_cm: float = 1.0,
        default_y_start_cm: float = 1.0,
        default_y_step_cm: float = 11.0,
):
    """
    Вставляет PNG из папки image_folder в лист sheet_name файла xlsx_path.
    images_config — словарь вида:
    {
        "01.png": {"x_cm": 1.0, "y_cm": 50.0, "w_cm": 8.0, "h_cm": 6.0},
        "02.png": {"x_cm": 10.0, "y_cm": 50.0, "w_cm": 8.0, "h_cm": 6.0},
        "03.png": {"x_cm": 1.0, "y_cm": 57.0, "w_cm": 8.0, "h_cm": 6.0},
    }
    Если для картинки нет конфигурации — используются default_* параметры.
    """
    xlsx_path = Path(xlsx_path)
    image_folder = Path(image_folder)
    png_files = sorted(image_folder.glob("*.png"))
    wb = load_workbook(xlsx_path)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Лист '{sheet_name}' не найден. Доступные: {wb.sheetnames}")
    ws = wb[sheet_name]
    # Если images_config не задан — используем автоматическую расстановку
    if images_config is None:
        images_config = {
            p.name: {
                "x_cm": default_x_cm,
                "y_cm": default_y_start_cm + i * default_y_step_cm,
                "w_cm": default_width_cm,
                "h_cm": default_height_cm,
            }
            for i, p in enumerate(png_files)
        }
    for png_file in png_files:
        png_name = png_file.name
        # Получаем конфигурацию для этой картинки (или используем значения по умолчанию)
        if png_name in images_config:
            config = images_config[png_name]
        else:
            # Если картинки нет в конфиге — пропускаем или используем дефолт
            print(f"  ⚠ {png_name} не найден в images_config, используется конфиг по умолчанию")
            config = {
                "x_cm": default_x_cm,
                "y_cm": default_y_start_cm,
                "w_cm": default_width_cm,
                "h_cm": default_height_cm,
            }
        img = Image(str(png_file))
        width_emu = int(cm_to_EMU(config["w_cm"]))
        height_emu = int(cm_to_EMU(config["h_cm"]))
        ext = XDRPositiveSize2D(cx=width_emu, cy=height_emu)
        marker = AnchorMarker(
            col=0,
            colOff=int(cm_to_EMU(config["x_cm"])),
            row=0,
            rowOff=int(cm_to_EMU(config["y_cm"])),
        )
        anchor = OneCellAnchor(_from=marker, ext=ext)
        img.anchor = anchor
        ws.add_image(img)
    wb.save(xlsx_path)
    # print(f"\n✅ Сохранено: {xlsx_path.resolve()}")

# pdf ===========================
def convert_xlsx_to_pdf(xlsx_path: str, output_folder: str) -> str:
    """
    Конвертирует XLSX в PDF с помощью LibreOffice.
    Исправлена ошибка "Can't open display".
    """
    abs_xlsx_path = os.path.abspath(xlsx_path)
    abs_output_folder = os.path.abspath(output_folder)

    if not os.path.exists(abs_xlsx_path):
        raise FileNotFoundError(f"Исходный файл не найден: {abs_xlsx_path}")
    os.makedirs(abs_output_folder, exist_ok=True)
    unique_profile = f"/tmp/lo_profile_{os.getpid()}_{os.geteuid()}"
    command = [
        'libreoffice',
        '--headless',
        '--invisible',
        '--norestore',
        '--nofirststartwizard',
        '--convert-to', 'pdf',
        '--outdir', abs_output_folder,
        f'-env:UserInstallation=file://{unique_profile}',
        abs_xlsx_path
    ]
    env = os.environ.copy()
    # Очищаем от мусора cv2/OpenCV
    vars_to_remove = ['LD_LIBRARY_PATH', 'QT_PLUGIN_PATH', 'QT_QPA_PLATFORM_PLUGIN_PATH', 'PYTHONPATH']
    for var in vars_to_remove:
        env.pop(var, None)
    # ИСПРАВЛЕНИЕ: Устанавливаем фиктивный дисплей вместо удаления
    env['DISPLAY'] = ':0'
    env['SAL_USE_VCLPLUGIN'] = 'gen'
    env['QT_QPA_PLATFORM'] = 'offscreen'
    try:
        result = subprocess.run(
            command,
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            env=env,
            timeout=60
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError(f"LibreOffice завис при конвертации: {abs_xlsx_path}")
    except subprocess.CalledProcessError as e:
        raise RuntimeError(f"Ошибка LibreOffice:\nSTDOUT:\n{e.stdout}\nSTDERR:\n{e.stderr}")
    except FileNotFoundError:
        raise RuntimeError("LibreOffice не найден. Установите: sudo apt install libreoffice")

    input_filename = Path(abs_xlsx_path).stem
    pdf_filename = f"{input_filename}.pdf"
    pdf_path = Path(abs_output_folder) / pdf_filename
    if not pdf_path.exists():
        fallback_path = Path.cwd() / pdf_filename
        if fallback_path.exists():
            import shutil
            shutil.move(str(fallback_path), str(pdf_path))
        else:
            files_in_output = list(Path(abs_output_folder).glob("*"))
            raise RuntimeError(
                f"Файл PDF не был создан.\n"
                f"Ожидался: {pdf_path}\n"
                f"Файлы в папке: {files_in_output}"
            )
    return str(pdf_path)


if __name__ == "__main__":
    pass

